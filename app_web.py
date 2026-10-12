from flask import Flask, request, render_template_string, send_file
from werkzeug.utils import secure_filename
from pathlib import Path
from io import BytesIO
import subprocess
import tempfile
import sys
import json
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
CONVERSOR = BASE_DIR / "conversor.py"

HTML = """
<!doctype html>
<html lang="pt-br">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Conversor CSV</title>
<style>
* { box-sizing: border-box; }
body {
    font-family: Arial, sans-serif;
    background: #f1f4f8;
    margin: 0;
    padding: 20px;
    color: #202938;
}
main {
    max-width: 560px;
    margin: 25px auto;
    background: white;
    padding: 25px;
    border-radius: 12px;
    box-shadow: 0 3px 12px #00000015;
}
h1 { font-size: 25px; margin-top: 0; }
p { line-height: 1.5; }
label { display: block; margin: 18px 0 8px; font-weight: bold; }
input[type=file], select {
    width: 100%;
    padding: 12px;
    border: 1px solid #aeb8c5;
    border-radius: 6px;
    background: white;
    font-size: 16px;
}
button {
    width: 100%;
    margin-top: 22px;
    padding: 14px;
    border: 0;
    border-radius: 6px;
    background: #1769aa;
    color: white;
    font-size: 17px;
    font-weight: bold;
}
.aviso {
    margin-top: 18px;
    padding: 12px;
    border-radius: 6px;
    background: #fff0f0;
    color: #9b1c1c;
    overflow-wrap: anywhere;
}
.info { color: #526174; font-size: 14px; }
</style>
</head>
<body>
<main>
<h1>Conversor CSV para JSON ou Excel</h1>
<p>Selecione um arquivo CSV e escolha o formato de saída.</p>

<form method="post" enctype="multipart/form-data">
    <label for="arquivo">Arquivo CSV</label>
    <input type="file" id="arquivo" name="arquivo"
           accept=".csv,text/csv" required>

    <label for="formato">Formato de saída</label>
    <select id="formato" name="formato">
        <option value="xlsx" selected>Excel (.xlsx) — planilha formatada</option>
        <option value="json">JSON (.json) — dados para sistemas</option>
    </select>

    <button type="submit">Converter e baixar arquivo</button>
</form>

{% if mensagem %}
<div class="aviso"><strong>Não foi possível concluir:</strong><br>
{{ mensagem }}</div>
{% endif %}

<p class="info">
O CSV deve ter cabeçalho. Arquivos de até 5 MB são aceitos.
</p>
</main>
</body>
</html>
"""


def pagina(mensagem=None, status=200):
    return render_template_string(HTML, mensagem=mensagem), status


def criar_excel(registros, cabecalhos):
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "Dados"
    planilha.freeze_panes = "A2"
    planilha.sheet_view.showGridLines = False

    cor_cabecalho = "174A70"
    cor_linha = "EAF2F8"
    cor_borda = "D5DEE7"

    preenchimento_cabecalho = PatternFill(
        fill_type="solid", fgColor=cor_cabecalho
    )
    preenchimento_alternado = PatternFill(
        fill_type="solid", fgColor=cor_linha
    )
    borda_fina = Border(
        bottom=Side(style="thin", color=cor_borda)
    )

    for coluna, nome in enumerate(cabecalhos, start=1):
        celula = planilha.cell(row=1, column=coluna, value=nome)
        celula.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
        celula.fill = preenchimento_cabecalho
        celula.alignment = Alignment(
            horizontal="left", vertical="center"
        )
        celula.border = borda_fina

    planilha.row_dimensions[1].height = 30

    for numero_linha, registro in enumerate(registros, start=2):
        for coluna, nome in enumerate(cabecalhos, start=1):
            valor = registro.get(nome, "")
            celula = planilha.cell(
                row=numero_linha, column=coluna, value=valor
            )
            celula.font = Font(name="Arial", size=13)
            celula.alignment = Alignment(
                vertical="center", wrap_text=True
            )
            celula.border = borda_fina

            if numero_linha % 2 == 0:
                celula.fill = preenchimento_alternado

        planilha.row_dimensions[numero_linha].height = 24

    for coluna, nome in enumerate(cabecalhos, start=1):
        maior = len(str(nome))

        for registro in registros:
            valor = registro.get(nome, "")
            maior = max(maior, len(str(valor)))

        planilha.column_dimensions[
            get_column_letter(coluna)
        ].width = min(max(maior + 4, 16), 35)

    if cabecalhos:
        planilha.auto_filter.ref = (
            f"A1:{get_column_letter(len(cabecalhos))}"
            f"{max(len(registros) + 1, 1)}"
        )

    planilha.sheet_properties.pageSetUpPr.fitToPage = True
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 0
    planilha.page_setup.orientation = "landscape"

    arquivo = BytesIO()
    pasta.save(arquivo)
    arquivo.seek(0)
    return arquivo


@app.route("/", methods=["GET", "POST"])
def inicio():
    if request.method == "GET":
        return pagina()[0]

    arquivo = request.files.get("arquivo")
    formato = request.form.get("formato", "xlsx")

    if formato not in ("json", "xlsx"):
        return pagina("Formato de saída inválido.", 400)

    if not arquivo or not arquivo.filename:
        return pagina("Selecione um arquivo CSV.", 400)

    nome_seguro = secure_filename(arquivo.filename)

    if not nome_seguro or Path(nome_seguro).suffix.lower() != ".csv":
        return pagina("Selecione um arquivo com extensão .csv.", 400)

    if not CONVERSOR.is_file():
        return pagina("O arquivo conversor.py não foi encontrado.", 500)

    try:
        with tempfile.TemporaryDirectory() as temporario:
            pasta_temp = Path(temporario)
            entrada = pasta_temp / nome_seguro
            saida_json = pasta_temp / "resultado.json"

            arquivo.save(entrada)

            resultado = subprocess.run(
                [
                    sys.executable,
                    str(CONVERSOR),
                    str(entrada),
                    str(saida_json),
                ],
                cwd=str(BASE_DIR),
                capture_output=True,
                text=True,
                timeout=30,
            )

            if resultado.returncode != 0 or not saida_json.is_file():
                mensagem = (
                    resultado.stderr.strip()
                    or resultado.stdout.strip()
                    or "O CSV não passou na validação."
                )
                return pagina(mensagem, 400)

            if formato == "json":
                conteudo_json = saida_json.read_bytes()
                arquivo_memoria = BytesIO(conteudo_json)
                return send_file(
                    arquivo_memoria,
                    mimetype="application/json",
                    as_attachment=True,
                    download_name=f"{Path(nome_seguro).stem}.json",
                )

            registros = json.loads(
                saida_json.read_text(encoding="utf-8")
            )

            # Recupera os cabeçalhos mesmo quando o CSV não tem registros.
            with entrada.open(
                "r", encoding="utf-8-sig", newline=""
            ) as arquivo_csv:
                leitor = csv.reader(arquivo_csv)
                cabecalhos = next(leitor, [])

            planilha = criar_excel(registros, cabecalhos)

            return send_file(
                planilha,
                mimetype=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                as_attachment=True,
                download_name=f"{Path(nome_seguro).stem}.xlsx",
            )

    except subprocess.TimeoutExpired:
        return pagina("A conversão demorou demais e foi interrompida.", 500)
    except (OSError, UnicodeError, ValueError, csv.Error) as erro:
        return pagina(f"Erro ao processar o arquivo: {erro}", 400)
    except Exception:
        return pagina(
            "Ocorreu um erro interno ao processar o arquivo.",
            500,
        )


@app.errorhandler(413)
def arquivo_grande(erro):
    return pagina("O arquivo ultrapassa o limite de 5 MB.", 413)


if __name__ == "__main__":
    print("Conversor web iniciado.")
    print("Abra no navegador: http://127.0.0.1:8000")
    app.run(host="127.0.0.1", port=8000, debug=False)
