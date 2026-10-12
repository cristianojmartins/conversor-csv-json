import io
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app_web import app


class TestAplicacaoWeb(unittest.TestCase):
    def setUp(self):
        app.config["TESTING"] = True
        self.cliente = app.test_client()

    def enviar_csv(self, conteudo, formato="json",
                   nome="clientes.csv"):
        return self.cliente.post(
            "/",
            data={
                "arquivo": (
                    io.BytesIO(conteudo),
                    nome,
                ),
                "formato": formato,
            },
            content_type="multipart/form-data",
        )

    def test_pagina_inicial(self):
        resposta = self.cliente.get("/")
        self.assertEqual(resposta.status_code, 200)

    def test_conversao_json_com_acentos(self):
        csv = (
            "nome,cidade,idade\n"
            "João,Itaperuna,30\n"
            "Maria,Niterói,25\n"
        ).encode("utf-8")

        resposta = self.enviar_csv(csv, "json")

        self.assertEqual(resposta.status_code, 200)
        self.assertIn("application/json", resposta.content_type)
        self.assertIn("João".encode(), resposta.data)
        self.assertIn("Itaperuna".encode(), resposta.data)
        self.assertTrue(resposta.headers.get("Content-Disposition"))

    def test_conversao_excel(self):
        csv = (
            "nome,cidade,idade\n"
            "João,Itaperuna,30\n"
            "Maria,Niterói,25\n"
        ).encode("utf-8")

        resposta = self.enviar_csv(csv, "xlsx")

        self.assertEqual(resposta.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            resposta.content_type,
        )

        arquivo = load_workbook(
            io.BytesIO(resposta.data),
            read_only=True,
            data_only=True,
        )
        planilha = arquivo.active
        linhas = list(planilha.iter_rows(values_only=True))

        self.assertEqual(linhas[0], ("nome", "cidade", "idade"))
        self.assertEqual(linhas[1], ("João", "Itaperuna", 30))
        self.assertEqual(linhas[2], ("Maria", "Niterói", 25))
        arquivo.close()

    def test_rejeita_extensao_nao_permitida(self):
        resposta = self.enviar_csv(
            b"nome,idade\nAna,20\n",
            "json",
            "clientes.txt",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_idade_invalida(self):
        resposta = self.enviar_csv(
            b"nome,idade\nAna,abc\n",
            "json",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_cabecalho_duplicado(self):
        resposta = self.enviar_csv(
            b"nome,nome\nAna,Joana\n",
            "json",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_csv_sem_cabecalho(self):
        resposta = self.enviar_csv(b"", "json")
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_formato_nao_suportado(self):
        resposta = self.enviar_csv(
            b"nome,idade\nAna,20\n",
            "pdf",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_campos_ausentes(self):
        resposta = self.enviar_csv(
            b"nome,cidade,idade\nAna,Itaperuna\n",
            "json",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_campos_extras(self):
        resposta = self.enviar_csv(
            b"nome,idade\nAna,20,extra\n",
            "json",
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_envio_sem_arquivo(self):
        resposta = self.cliente.post(
            "/",
            data={"formato": "json"},
        )
        self.assertEqual(resposta.status_code, 400)

    def test_rejeita_arquivo_maior_que_limite(self):
        csv_grande = b"a" * (5 * 1024 * 1024 + 1)
        resposta = self.enviar_csv(csv_grande, "json")
        self.assertEqual(resposta.status_code, 413)

    def test_nome_de_arquivo_seguro(self):
        csv = b"nome,idade\nAna,20\n"

        resposta = self.enviar_csv(
            csv,
            "json",
            "../../clientes.csv",
        )

        self.assertEqual(resposta.status_code, 200)
        disposicao = resposta.headers.get(
            "Content-Disposition", ""
        )
        self.assertNotIn("../../", disposicao)

    def test_erro_interno_nao_expoe_detalhes(self):
        csv = b"nome,idade\nAna,20\n"

        with patch(
            "app_web.subprocess.run",
            side_effect=RuntimeError("ERRO_INTERNO_CONFIDENCIAL"),
        ):
            resposta = self.enviar_csv(csv, "json")

        self.assertEqual(resposta.status_code, 500)
        self.assertNotIn(
            b"ERRO_INTERNO_CONFIDENCIAL",
            resposta.data,
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
